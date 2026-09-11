"""Geração do relatório mensal executivo em PowerPoint."""

from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


LARGURA = 13.333
ALTURA = 7.5
MARGEM = 0.58
FONTE = "Aptos"
FUNDO = "F6F8FB"
TEXTO = "10243D"
MUTED = "66788D"
LINHA = "DCE5EF"


IDENTIDADES = {
    "ERIMAX": {
        "primaria": "075EA8",
        "secundaria": "FF8200",
        "suave": "EAF5FC",
        "logo": "erimax-logo.png",
    },
    "MARK": {
        "primaria": "0D5C93",
        "secundaria": "C61E2D",
        "suave": "EFF5FA",
        "logo": "mark-logo.png",
    },
}


def _rgb(cor):
    return RGBColor.from_string(cor)


def _moeda(valor):
    return f"R$ {float(valor or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _numero(valor, casas=0):
    return f"{float(valor or 0):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _logo(empresa):
    raiz = Path(__file__).resolve().parent
    nome = IDENTIDADES[empresa]["logo"]
    # O Flask publica ``public`` como diretório estático. Mantemos o fallback
    # para instalações locais antigas que ainda possuem a pasta ``static``.
    for pasta in (raiz / "public", raiz / "static"):
        caminho = pasta / nome
        if caminho.exists():
            return caminho
    return raiz / "public" / nome


def _caixa(slide, esquerda, topo, largura, altura, cor, raio=True):
    forma = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE,
        Inches(esquerda),
        Inches(topo),
        Inches(largura),
        Inches(altura),
    )
    forma.fill.solid()
    forma.fill.fore_color.rgb = _rgb(cor)
    forma.line.fill.background()
    return forma


def _texto(slide, texto, esquerda, topo, largura, altura, tamanho=18, cor=TEXTO, negrito=False, alinhamento=PP_ALIGN.LEFT):
    caixa = slide.shapes.add_textbox(
        Inches(esquerda), Inches(topo), Inches(largura), Inches(altura)
    )
    quadro = caixa.text_frame
    quadro.clear()
    quadro.word_wrap = True
    quadro.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragrafo = quadro.paragraphs[0]
    paragrafo.alignment = alinhamento
    trecho = paragrafo.add_run()
    trecho.text = texto
    fonte = trecho.font
    fonte.name = FONTE
    fonte.size = Pt(tamanho)
    fonte.bold = negrito
    fonte.color.rgb = _rgb(cor)
    return caixa


def _novo_slide(apresentacao, empresa, titulo=None, subtitulo=None):
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(FUNDO)
    identidade = IDENTIDADES[empresa]
    _caixa(slide, 0, 0, LARGURA, 0.12, identidade["secundaria"], raio=False)
    if titulo:
        _texto(slide, titulo, MARGEM, 0.42, 9.6, 0.45, 25, TEXTO, True)
    if subtitulo:
        _texto(slide, subtitulo, MARGEM, 0.89, 9.8, 0.3, 10.5, MUTED)
    logo = _logo(empresa)
    if logo.exists():
        slide.shapes.add_picture(str(logo), Inches(11.55), Inches(0.32), width=Inches(1.18))
    _texto(slide, empresa, MARGEM, 7.08, 2.0, 0.18, 8, MUTED, True)
    return slide


def _adicionar_transicao_fade(slide):
    """Inclui a transição Fade padrão do formato Open XML.

    ``python-pptx`` ainda não expõe transições na API pública. O elemento
    ``p:fade`` é uma transição oficial e simples do PowerPoint; por isso é
    seguro adicioná-lo sem recorrer ao Morph, que depende de extensões de
    versão e pode falhar em outros leitores de PPTX.
    """
    raiz = slide._element
    for filho in list(raiz):
        if filho.tag == qn("p:transition"):
            raiz.remove(filho)

    transicao = OxmlElement("p:transition")
    transicao.set("spd", "med")
    transicao.set("advClick", "1")
    transicao.append(OxmlElement("p:fade"))

    indice = len(raiz)
    for posicao, filho in enumerate(raiz):
        if filho.tag in (qn("p:timing"), qn("p:extLst")):
            indice = posicao
            break
    raiz.insert(indice, transicao)


def _adicionar_transicao_morph(slide):
    """Adiciona Morph com fallback oficial para Fade.

    O PowerPoint armazena Morph como uma extensão Office 2015 dentro de
    ``mc:AlternateContent``. Programas que não conhecem essa extensão leem o
    fallback ``p:fade`` e continuam abrindo o arquivo normalmente.
    """
    raiz = slide._element
    for filho in list(raiz):
        if filho.tag == qn("p:transition"):
            raiz.remove(filho)

    xml = b"""
        <mc:AlternateContent
            xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
            xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
            <mc:Choice
                xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main"
                Requires="p159">
                <p:transition
                    xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"
                    spd="med"
                    p14:dur="650"
                    advClick="1">
                    <p159:morph option="byObject"/>
                </p:transition>
            </mc:Choice>
            <mc:Fallback>
                <p:transition spd="med" advClick="1">
                    <p:fade/>
                </p:transition>
            </mc:Fallback>
        </mc:AlternateContent>
    """
    conteudo_alternativo = etree.fromstring(xml)
    indice = len(raiz)
    for posicao, filho in enumerate(raiz):
        if filho.tag in (qn("p:timing"), qn("p:extLst")):
            indice = posicao
            break
    raiz.insert(indice, conteudo_alternativo)


def _card(slide, titulo, valor, esquerda, topo, largura, empresa, destaque=False):
    identidade = IDENTIDADES[empresa]
    cor = identidade["primaria"] if destaque else "FFFFFF"
    valor_cor = "FFFFFF" if destaque else TEXTO
    rotulo_cor = "DDEEFF" if destaque else MUTED
    _caixa(slide, esquerda, topo, largura, 1.22, cor)
    _texto(slide, titulo.upper(), esquerda + 0.2, topo + 0.16, largura - 0.4, 0.22, 8.5, rotulo_cor, True)
    _texto(slide, valor, esquerda + 0.2, topo + 0.48, largura - 0.4, 0.46, 19, valor_cor, True)


def _chart(slide, tipo, titulo, categorias, series, esquerda, topo, largura, altura, empresa, rotacionar=False):
    if not categorias or not series:
        return None
    dados = CategoryChartData()
    dados.categories = categorias
    for nome, valores in series:
        dados.add_series(nome, valores)
    grafico = slide.shapes.add_chart(
        tipo, Inches(esquerda), Inches(topo), Inches(largura), Inches(altura), dados
    ).chart
    grafico.has_legend = len(series) > 1
    if grafico.has_legend:
        grafico.legend.include_in_layout = False
    grafico.has_title = True
    grafico.chart_title.text_frame.text = titulo
    grafico.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONTE
    grafico.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    grafico.value_axis.has_major_gridlines = True
    grafico.value_axis.major_gridlines.format.line.color.rgb = _rgb(LINHA)
    grafico.value_axis.tick_labels.font.name = FONTE
    grafico.value_axis.tick_labels.font.size = Pt(8)
    grafico.category_axis.tick_labels.font.name = FONTE
    grafico.category_axis.tick_labels.font.size = Pt(8)
    if rotacionar:
        grafico.category_axis.tick_labels.font.size = Pt(7)
    paleta = [IDENTIDADES[empresa]["primaria"], IDENTIDADES[empresa]["secundaria"], "6F95B7"]
    for indice, serie in enumerate(grafico.series):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = _rgb(paleta[indice % len(paleta)])
        serie.format.line.color.rgb = _rgb(paleta[indice % len(paleta)])
    return grafico


def _tabela(slide, cabecalhos, linhas, esquerda, topo, largura, altura, empresa):
    if not linhas:
        return
    forma = slide.shapes.add_table(
        len(linhas) + 1, len(cabecalhos), Inches(esquerda), Inches(topo), Inches(largura), Inches(altura)
    )
    tabela = forma.table
    identidade = IDENTIDADES[empresa]
    for coluna, valor in enumerate(cabecalhos):
        celula = tabela.cell(0, coluna)
        celula.text = valor
        celula.fill.solid()
        celula.fill.fore_color.rgb = _rgb(identidade["primaria"])
    for linha, valores in enumerate(linhas, start=1):
        for coluna, valor in enumerate(valores):
            celula = tabela.cell(linha, coluna)
            celula.text = str(valor)
            celula.fill.solid()
            celula.fill.fore_color.rgb = _rgb("FFFFFF" if linha % 2 else identidade["suave"])
    for linha in tabela.rows:
        for celula in linha.cells:
            celula.margin_left = Inches(0.08)
            celula.margin_right = Inches(0.08)
            for paragrafo in celula.text_frame.paragraphs:
                paragrafo.alignment = PP_ALIGN.LEFT
                for trecho in paragrafo.runs:
                    trecho.font.name = FONTE
                    trecho.font.size = Pt(8.5)
                    trecho.font.color.rgb = _rgb("FFFFFF" if linha == tabela.rows[0] else TEXTO)
                    trecho.font.bold = linha == tabela.rows[0]


def _capa(apresentacao, empresa, relatorio):
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[6])
    identidade = IDENTIDADES[empresa]
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(identidade["primaria"])
    _caixa(slide, 0, 0, 0.18, ALTURA, identidade["secundaria"], raio=False)
    _texto(slide, "RELATÓRIO MENSAL", 1.02, 1.45, 7.8, 0.45, 17, "DDEEFF", True)
    _texto(slide, empresa, 1.02, 2.0, 8.2, 0.9, 39, "FFFFFF", True)
    _texto(slide, f"{relatorio['nome_mes']} de {relatorio['ano']}", 1.02, 3.03, 7.0, 0.48, 23, "FFFFFF")
    _texto(slide, "Controle de despesas de viagem", 1.02, 3.73, 6.8, 0.3, 13, "DDEEFF")
    _caixa(slide, 1.02, 4.35, 2.8, 0.58, identidade["secundaria"])
    _texto(slide, "VIA GESTÃO", 1.24, 4.47, 2.36, 0.2, 10, "FFFFFF", True, PP_ALIGN.CENTER)

    # O painel claro evita que o azul escuro da logo desapareça no fundo
    # institucional, tanto para ERIMAX quanto para MARK.
    _caixa(slide, 8.85, 1.22, 3.75, 1.82, "FFFFFF")
    _caixa(slide, 8.85, 1.22, 3.75, 0.1, identidade["secundaria"], raio=False)
    logo = _logo(empresa)
    if logo.exists():
        slide.shapes.add_picture(str(logo), Inches(9.16), Inches(1.58), width=Inches(3.12))
    _texto(slide, "ViaGestão", 1.02, 6.67, 2.0, 0.25, 10, "DDEEFF", True)


def gerar_relatorio_powerpoint(empresa, relatorio, transicao="morph"):
    """Cria um PPTX 16:9 em memória, com gráficos editáveis no PowerPoint."""
    empresa = empresa.upper()
    if empresa not in IDENTIDADES:
        raise ValueError("Empresa não suportada para apresentação.")

    apresentacao = Presentation()
    apresentacao.slide_width = Inches(LARGURA)
    apresentacao.slide_height = Inches(ALTURA)
    _capa(apresentacao, empresa, relatorio)

    dados = relatorio["dados"]
    despesas = float(dados["total_despesa_mes"])
    receita = float(dados["total_receita_mes"])
    resultado = relatorio["resultado"]
    viagens = relatorio["viagens"]
    cargas = relatorio["cargas"]

    resumo = _novo_slide(
        apresentacao,
        empresa,
        "Resumo executivo",
        f"Indicadores de {relatorio['nome_mes']} de {relatorio['ano']}",
    )
    _card(resumo, "Viagens iniciadas", str(len(viagens)), 0.58, 1.42, 2.85, empresa)
    _card(resumo, "Faturamento", _moeda(receita), 3.62, 1.42, 2.85, empresa, True)
    _card(resumo, "Despesas", _moeda(despesas), 6.67, 1.42, 2.85, empresa)
    _card(resumo, "Resultado", _moeda(resultado), 9.72, 1.42, 2.85, empresa, True)
    _card(resumo, "Quilômetros", f"{_numero(relatorio['km_total'])} km", 0.58, 3.08, 3.75, empresa)
    _card(resumo, "Cargas lançadas", str(len(cargas)), 4.78, 3.08, 3.75, empresa)
    media = receita / len(viagens) if viagens else 0.0
    _card(resumo, "Receita por viagem", _moeda(media), 8.98, 3.08, 3.75, empresa)
    _texto(
        resumo,
        "Os indicadores utilizam os mesmos lançamentos e filtros do relatório mensal exibido no ViaGestão.",
        0.58,
        5.24,
        10.7,
        0.35,
        11,
        MUTED,
    )

    financeiro = _novo_slide(
        apresentacao,
        empresa,
        "Visão financeira",
        "Receitas, despesas e resultado apurados no período selecionado",
    )
    _chart(
        financeiro,
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        "Resultado mensal",
        [relatorio["nome_mes"]],
        [("Receita", [receita]), ("Despesa", [despesas]), ("Resultado", [resultado])],
        0.65,
        1.35,
        7.25,
        4.9,
        empresa,
    )
    taxa = dados["percentual_despesa_sobre_receita"]
    _card(financeiro, "Faturamento", _moeda(receita), 8.45, 1.63, 3.85, empresa, True)
    _card(financeiro, "Despesa sobre receita", f"{_numero(taxa, 2)}%" if taxa is not None else "—", 8.45, 3.12, 3.85, empresa)
    _card(financeiro, "Resultado", _moeda(resultado), 8.45, 4.61, 3.85, empresa)

    if relatorio["categorias"]:
        despesas_slide = _novo_slide(
            apresentacao,
            empresa,
            "Despesas por categoria",
            "Principais despesas lançadas no período",
        )
        principais = relatorio["categorias"][:8]
        _chart(
            despesas_slide,
            XL_CHART_TYPE.BAR_CLUSTERED,
            "Total por categoria",
            [item["nome"] for item in principais],
            [("Despesas", [item["valor"] for item in principais])],
            0.65,
            1.32,
            7.3,
            4.95,
            empresa,
        )
        linhas = [
            [item["nome"], _moeda(item["valor"]), f"{_numero(item['percentual'], 1)}%"]
            for item in principais[:6]
        ]
        _tabela(despesas_slide, ["Categoria", "Valor", "Participação"], linhas, 8.3, 1.62, 4.25, 3.85, empresa)

    if cargas:
        cargas_slide = _novo_slide(
            apresentacao,
            empresa,
            "Cargas e empresas atendidas",
            "Receita de entrega e coleta registrada no período",
        )
        empresas_atendidas = list(dados["por_empresa"].items())[:8]
        _chart(
            cargas_slide,
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            "Receita por empresa atendida",
            [nome for nome, _ in empresas_atendidas],
            [("Entrega", [item["entrega"] for _, item in empresas_atendidas]), ("Coleta", [item["coleta"] for _, item in empresas_atendidas])],
            0.65,
            1.32,
            7.3,
            4.95,
            empresa,
            True,
        )
        _card(cargas_slide, "Entregas", _moeda(sum(item["entrega"] for _, item in empresas_atendidas)), 8.45, 1.62, 3.85, empresa, True)
        _card(cargas_slide, "Coletas", _moeda(sum(item["coleta"] for _, item in empresas_atendidas)), 8.45, 3.12, 3.85, empresa)
        _card(cargas_slide, "Lançamentos", str(len(cargas)), 8.45, 4.61, 3.85, empresa)

    if viagens:
        viagens_slide = _novo_slide(
            apresentacao,
            empresa,
            "Viagens do mês",
            "Rotas com maior receita entre as viagens iniciadas no período",
        )
        linhas = [
            [item["rota"], item["motorista"], f"{_numero(item['km'])} km", _moeda(item["receita"]), _moeda(item["despesa"])]
            for item in relatorio["viagens_resumo"][:7]
        ]
        _tabela(
            viagens_slide,
            ["Rota", "Motorista", "KM", "Receita", "Despesa"],
            linhas,
            0.65,
            1.37,
            12.05,
            4.85,
            empresa,
        )
        _texto(viagens_slide, f"{len(viagens)} viagem(ns) iniciada(s) no período", 0.65, 6.36, 4.3, 0.25, 10.5, MUTED)

    if relatorio["motoristas"]:
        motoristas_slide = _novo_slide(
            apresentacao,
            empresa,
            "Desempenho por motorista",
            "Comparativo de receita nas viagens iniciadas no período",
        )
        principais = relatorio["motoristas"][:8]
        _chart(
            motoristas_slide,
            XL_CHART_TYPE.BAR_CLUSTERED,
            "Receita por motorista",
            [item["motorista"] for item in principais],
            [("Receita", [item["receita"] for item in principais])],
            0.65,
            1.32,
            7.3,
            4.95,
            empresa,
        )
        linhas = [
            [item["motorista"], item["viagens"], f"{_numero(item['km'])} km", _moeda(item["receita"])]
            for item in principais[:6]
        ]
        _tabela(motoristas_slide, ["Motorista", "Viagens", "KM", "Receita"], linhas, 8.3, 1.62, 4.25, 3.85, empresa)

    if relatorio["rotas"]:
        rotas_slide = _novo_slide(
            apresentacao,
            empresa,
            "Principais rotas",
            "Consolidação por origem e destino",
        )
        linhas = [
            [item["rota"], item["viagens"], f"{_numero(item['km'])} km", _moeda(item["receita"]), _moeda(item["despesa"])]
            for item in relatorio["rotas"][:8]
        ]
        _tabela(rotas_slide, ["Rota", "Viagens", "KM", "Receita", "Despesa"], linhas, 0.65, 1.37, 12.05, 4.85, empresa)

    historico = relatorio["historico"]
    if any(item["receita"] or item["despesa"] for item in historico):
        historico_slide = _novo_slide(
            apresentacao,
            empresa,
            "Evolução mensal",
            "Receita e despesas dos últimos seis meses",
        )
        _chart(
            historico_slide,
            XL_CHART_TYPE.LINE_MARKERS,
            "Histórico financeiro",
            [item["rotulo"] for item in historico],
            [("Receita", [item["receita"] for item in historico]), ("Despesa", [item["despesa"] for item in historico])],
            0.75,
            1.35,
            11.8,
            4.95,
            empresa,
        )

    encerramento = apresentacao.slides.add_slide(apresentacao.slide_layouts[6])
    identidade = IDENTIDADES[empresa]
    encerramento.background.fill.solid()
    encerramento.background.fill.fore_color.rgb = _rgb(identidade["primaria"])
    _caixa(encerramento, 0, 0, 0.18, ALTURA, identidade["secundaria"], raio=False)
    _texto(encerramento, "RELATÓRIO MENSAL", 1.05, 1.72, 7.3, 0.32, 15, "DDEEFF", True)
    _texto(encerramento, empresa, 1.05, 2.18, 7.3, 0.62, 31, "FFFFFF", True)
    _texto(encerramento, f"{relatorio['nome_mes']} de {relatorio['ano']}", 1.05, 3.02, 7.3, 0.35, 17, "FFFFFF")
    _texto(encerramento, "ViaGestão", 1.05, 5.85, 2.0, 0.25, 10, "DDEEFF", True)

    _caixa(encerramento, 8.85, 1.62, 3.75, 1.82, "FFFFFF")
    _caixa(encerramento, 8.85, 1.62, 3.75, 0.1, identidade["secundaria"], raio=False)
    logo = _logo(empresa)
    if logo.exists():
        encerramento.shapes.add_picture(str(logo), Inches(9.16), Inches(1.98), width=Inches(3.12))

    for indice, slide in enumerate(apresentacao.slides):
        if transicao == "morph" and indice > 0:
            _adicionar_transicao_morph(slide)
        else:
            _adicionar_transicao_fade(slide)

    arquivo = BytesIO()
    apresentacao.save(arquivo)
    arquivo.seek(0)
    return arquivo
